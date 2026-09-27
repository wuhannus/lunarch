#!/usr/bin/env python3
"""Collect the worst setup (critical) path of every run into one Excel."""
import re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = '/mnt/d/opencode/lunarch/flow'
RUNS = [
    ('baseline (7.0 ns, delay ABC)',      7.0,  0, f'{BASE}/baseline/6_finish.rpt'),
    ('timing run (5.0 ns, delay ABC)',    5.0,  0, f'{BASE}/opt/timing_diematch/6_finish.rpt'),
    ('power run (20.0 ns, area ABC)',    20.0,  1, f'{BASE}/opt/power_diematch/6_finish.rpt'),
    ('sweep p7_delay (7.0 ns, delay)',    7.0,  0, f'{BASE}/opt/sweep/p7_delay.rpt'),
    ('sweep p10_delay (10.0 ns, delay)', 10.0,  0, f'{BASE}/opt/sweep/p10_delay.rpt'),
    ('sweep p7_area (7.0 ns, area)',      7.0,  1, f'{BASE}/opt/sweep/p7_area.rpt'),
    ('sweep p10_area (10.0 ns, area)',   10.0,  1, f'{BASE}/opt/sweep/p10_area.rpt'),
]
OUT = f'{BASE}/opt/critical_paths.xlsx'


def num(pat, txt, grp=1):
    m = re.search(pat, txt, re.M)
    return float(m.group(grp)) if m else None


def parse(path):
    if not os.path.exists(path):
        return None
    t = open(path, encoding='utf-8', errors='ignore').read()
    d = {}
    d['wns'] = num(r'^wns max\s+(-?[\d.]+)', t)
    d['tns'] = num(r'^tns max\s+(-?[\d.]+)', t)
    d['pmin'] = num(r'period_min = ([\d.]+)', t)
    d['fmax'] = num(r'fmax = ([\d.]+)', t)
    m = re.search(r'^Total\s+[\d.eE+-]+\s+[\d.eE+-]+\s+[\d.eE+-]+\s+([\d.eE+-]+)', t, re.M)
    d['power'] = float(m.group(1)) * 1000 if m else None   # W -> mW
    # worst setup path (first in report_checks -path_delay max)
    sec = t.split('report_checks -path_delay max', 1)
    if len(sec) > 1:
        p = sec[1][:30000]
        ms = re.search(r'Startpoint:\s*(\S+)', p)
        me = re.search(r'Endpoint:\s*(\S+)', p)
        mg = re.search(r'Path Group:\s*(\S+)', p)
        ma = re.search(r'^[\s]*(-?[\d.]+)\s+data arrival time', p, re.M)
        mr = re.search(r'^[\s]*(-?[\d.]+)\s+data required time', p, re.M)
        msl = re.search(r'^[\s]*(-?[\d.]+)\s+slack', p, re.M)
        d['start'] = ms.group(1) if ms else '?'
        d['end'] = me.group(1) if me else '?'
        d['group'] = mg.group(1) if mg else '?'
        d['arrival'] = float(ma.group(1)) if ma else None
        d['required'] = float(mr.group(1)) if mr else None
        d['slack'] = float(msl.group(1)) if msl else None
    return d


HDR = Font(bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor='1F3A5F')
SUB = Font(bold=True)
BORD = Border(*[Side(style='thin', color='BBBBBB')] * 4)
ALT = PatternFill('solid', fgColor='F2F6FC')

wb = Workbook()
ws = wb.active
ws.title = 'Critical paths'
cols = ['Run', 'clk (ns)', 'ABC', 'Path group', 'Start point', 'End point',
        'Data arrival (ns)', 'Data required (ns)', 'Slack / WNS (ns)',
        'WNS (ns)', 'TNS (ns)', 'Min period (ns)', 'fmax (MHz)', 'Power (mW)']
ws.append(cols)
for c in range(1, len(cols) + 1):
    cell = ws.cell(row=1, column=c); cell.font = HDR; cell.fill = HF
    cell.alignment = Alignment(horizontal='center', vertical='center')
for name, clk, abc, path in RUNS:
    d = parse(path)
    if not d:
        ws.append([name, clk, 'delay' if abc == 0 else 'area'] + ['(report missing)'] * (len(cols) - 3))
        continue
    ws.append([name, clk, 'delay' if abc == 0 else 'area', d.get('group'),
               d.get('start'), d.get('end'), d.get('arrival'), d.get('required'),
               d.get('slack'), d.get('wns'), d.get('tns'), d.get('pmin'),
               d.get('fmax'), round(d['power'], 2) if d.get('power') else None])
widths = [30, 9, 7, 10, 34, 34, 16, 16, 15, 10, 11, 14, 11, 11]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.border = BORD
        if isinstance(c.value, float):
            c.number_format = '0.0000'
    if row[0].row % 2 == 0:
        for c in row:
            c.fill = ALT
ws.freeze_panes = 'A2'

# Notes sheet
n = wb.create_sheet('Notes')
for r in [
    ['Field', 'Meaning'],
    ['Run', 'flow configuration (clock constraint, ABC objective)'],
    ['clk (ns)', 'SDC clock-period constraint'],
    ['ABC', 'ABC_AREA: delay = 0 (speed), area = 1 (area-oriented mapping)'],
    ['Path group', 'STA clock group of the worst path (clk = core, vclk_clk = output)'],
    ['Start / End point', 'worst setup (critical) path endpoints'],
    ['Slack / WNS', 'slack of that worst path (= block setup WNS)'],
    ['Min period / fmax', 'report_clock_min_period (reg\u2192reg achievable)'],
    ['Power (mW)', 'post-route total power'],
    ['All runs', 'share the identical baseline die (376.625 \u00b5m)'],
]:
    n.append(r)
n['A1'].font = HDR; n['A1'].fill = HF; n['B1'].font = HDR; n['B1'].fill = HF
n.column_dimensions['A'].width = 22
n.column_dimensions['B'].width = 78

wb.save(OUT)
print('wrote', OUT)
for name, clk, abc, path in RUNS:
    d = parse(path)
    if d:
        print(f"  {name:34} start={d.get('start')} end={d.get('end')} slack={d.get('slack')}")
