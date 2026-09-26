#!/usr/bin/env python3
"""Parse an OpenROAD report_checks dump into a timing-summary Excel workbook."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = '/mnt/d/opencode/timing_full.rpt'
OUT = '/mnt/d/opencode/lunarch/flow/timing/riscv32i_timing_summary.xlsx'
PERIOD = 7.0

text = open(SRC, encoding='utf-8', errors='ignore').read()
blocks = re.split(r'\n(?=Startpoint:)', text)


def ltime(line):
    m = re.match(r'\s*(-?\d+\.\d+)\s+(-?\d+\.\d+)\s', line)
    return float(m.group(2)) if m else None


def func_group(start, end, body):
    b = body
    if end.startswith('aluout'):
        return 'PG8 result-out'
    if end.startswith('writedata'):
        return 'PG4 store-out'
    if end.startswith('pc[') or end == 'pc':
        return 'PG1 pc-out'
    if 'shifter' in b:
        return 'PG3 shift/write-back'
    if 'magcompare' in b:
        return 'PG5 branch-compare'
    if 'pcreg' in start and 'pcreg' in end:
        return 'PG1 fetch/next-PC'
    if 'pcreg' in end:
        return 'PG1/PG5 PC-redirect'
    if 'rf.rf' in start and 'rf.rf' in end:
        return 'PG2 ALU/write-back'
    if start.startswith('instr') or start.startswith('readdata'):
        return 'input->reg'
    return 'other'


rows = []
for b in blocks:
    b = b.strip('\n')
    if not b.startswith('Startpoint:'):
        continue
    lines = b.splitlines()
    start = lines[0].split('Startpoint:', 1)[1].strip()
    m = re.search(r'Endpoint:\s*(\S+)', b)
    end = m.group(1) if m else '?'
    m = re.search(r'Path Group:\s*(\S+)', b)
    group = m.group(1) if m else '?'
    m = re.search(r'Path Type:\s*(\S+)', b)
    ptype = m.group(1) if m else '?'
    m = re.search(r'^\s*([\d.]+)\s+data arrival time', b, re.M)
    arrival = float(m.group(1)) if m else None
    m = re.search(r'^\s*([\d.]+)\s+data required time', b, re.M)
    required = float(m.group(1)) if m else None
    m = re.search(r'^\s*(-?[\d.]+)\s+slack', b, re.M)
    slack = float(m.group(1)) if m else None

    arr_i = next((i for i, l in enumerate(lines) if 'data arrival time' in l), None)
    launch_i = max((i for i in range(arr_i) if '/CLK (' in lines[i]), default=None)
    launch = ltime(lines[launch_i]) if launch_i is not None else None

    cap_i = None
    for i in range(arr_i + 1, len(lines)):
        if '/CLK (' in lines[i]:
            cap_i = i
    if cap_i is not None:
        edge_i = next((i for i in range(arr_i + 1, len(lines)) if 'rise edge' in lines[i]), None)
        edge = ltime(lines[edge_i]) if edge_i is not None else PERIOD
        capture = ltime(lines[cap_i]) - edge
    else:
        edge_i = next((i for i in range(arr_i + 1, len(lines)) if 'rise edge' in lines[i]), None)
        edge = ltime(lines[edge_i]) if edge_i is not None else 0.0
        cn = next((i for i in range(arr_i + 1, len(lines)) if 'clock network delay' in lines[i]), None)
        capture = None
        if cn is not None:
            mm = re.match(r'\s*(-?\d+\.\d+)\s+(-?\d+\.\d+)\s', lines[cn])
            if mm:
                capture = float(mm.group(2)) - edge

    skew = (capture - launch) if (capture is not None and launch is not None) else None
    data_delay = (arrival - launch) if (arrival is not None and launch is not None) else None

    q_i = max((i for i in range(arr_i) if '/Q (' in lines[i]), default=None)
    levels = None
    if q_i is not None:
        e_i = max((i for i in range(q_i + 1, arr_i)
                   if '(sky130_fd_sc_hd__' in lines[i] or '(out)' in lines[i]), default=None)
        if e_i is not None:
            levels = sum(1 for i in range(q_i + 1, e_i) if '(sky130_fd_sc_hd__' in lines[i])

    rows.append(dict(start=start, end=end, group=group, type=ptype,
                     fgroup=func_group(start, end, b),
                     launch=launch, data=data_delay, levels=levels,
                     arrival=arrival, capture=capture, skew=skew,
                     required=required, slack=slack))

# ---------------- workbook ----------------
wb = Workbook()
HDR = Font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='3B6FB0')
SUB = Font(bold=True)
thin = Side(style='thin', color='BBBBBB')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR
        cell.fill = HFILL
        cell.alignment = Alignment(horizontal='center', vertical='center')


ws = wb.active
ws.title = 'Timing Paths'
headers = ['#', 'Type', 'Report group', 'Functional path group',
           'Start point', 'End point',
           'Launch clk latency (ns)', 'Data-path delay (ns)', 'Data-path cells (logic levels)',
           'Data arrival (ns)', 'Capture clk latency (ns)', 'Clock skew (ns)',
           'Data required (ns)', 'Slack / WNS (ns)', 'Status']
ws.append(headers)
for i, r in enumerate(rows, 1):
    status = 'MET' if (r['slack'] is not None and r['slack'] >= 0) else 'VIOLATED'
    ws.append([i, 'Setup (max)' if r['type'] == 'max' else 'Hold (min)',
               r['group'], r['fgroup'], r['start'], r['end'],
               r['launch'], r['data'], r['levels'], r['arrival'],
               r['capture'], r['skew'], r['required'], r['slack'], status])
style_header(ws, len(headers))
widths = [4, 12, 12, 24, 34, 34, 14, 14, 16, 12, 14, 12, 14, 14, 10]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.border = BORD
        if isinstance(cell.value, float):
            cell.number_format = '0.000'
    if row[0].row % 2 == 0:
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='F2F6FC')
ws.freeze_panes = 'A2'

# ---- Summary sheet ----
ws2 = wb.create_sheet('Summary')
ws2.append(['Baseline APR timing summary — sky130hd/riscv32i (riscv)'])
ws2['A1'].font = Font(bold=True, size=13)
ws2.append([])
ws2.append(['Clock period (ns)', PERIOD])
ws2['A4'] = 'Constraint / clock'
ws2['B4'], ws2['C4'], ws2['D4'] = 'clk period (ns)', 'fmax (MHz)', 'WNS (ns)'
for c in 'ABCD':
    ws2[f'{c}4'].font = SUB
setup = [r for r in rows if r['type'] == 'max']
hold = [r for r in rows if r['type'] == 'min']
wns = min((r['slack'] for r in setup if r['slack'] is not None), default=None)
tns = sum(r['slack'] for r in setup if r['slack'] is not None and r['slack'] < 0)
worst_hold = min((r['slack'] for r in hold if r['slack'] is not None), default=None)
ws2['A5'], ws2['B5'], ws2['C5'], ws2['D5'] = 'setup', PERIOD, round(1000.0 / PERIOD, 2), wns
ws2.append([])
ws2.append(['', 'Group', 'Worst slack (ns)', 'Worst start', 'Worst end', 'Paths'])
for c in range(1, 7):
    ws2.cell(row=ws2.max_row, column=c).font = SUB
groups = []
for g in sorted({r['group'] for r in setup}):
    gr = [r for r in setup if r['group'] == g]
    worst = min(gr, key=lambda r: (r['slack'] if r['slack'] is not None else 1e9))
    groups.append((g, worst['slack'], worst['start'], worst['end'], len(gr)))
for g in groups:
    ws2.append(['', g[0], g[1], g[2], g[3], g[4]])
ws2.append([])
ws2.append(['Total setup TNS (ns)', tns])
ws2.append(['Total setup WNS (ns)', wns])
ws2.append(['Worst hold slack (ns)', worst_hold])
ws2.column_dimensions['A'].width = 22
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 34
ws2.column_dimensions['D'].width = 34
ws2.column_dimensions['E'].width = 10
ws2.column_dimensions['F'].width = 8

# ---- Clock skew sheet ----
ws3 = wb.create_sheet('Clock Skew')
ws3.append(['Per-path clock skew (capture clk latency − launch clk latency)'])
ws3['A1'].font = Font(bold=True, size=12)
ws3.append([])
ws3.append(['Type', 'Report group', 'Functional path group', 'Start', 'End',
            'Launch clk lat (ns)', 'Capture clk lat (ns)', 'Clock skew (ns)'])
style_header(ws3, 8)
for r in rows:
    ws3.append(['Setup' if r['type'] == 'max' else 'Hold', r['group'], r['fgroup'],
                r['start'], r['end'], r['launch'], r['capture'], r['skew']])
for i, w in enumerate([8, 12, 24, 34, 34, 16, 16, 14], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
for row in ws3.iter_rows(min_row=4):
    for cell in row:
        cell.border = BORD
        if isinstance(cell.value, float):
            cell.number_format = '0.0000'

# ---- Notes sheet ----
ws4 = wb.create_sheet('Notes')
notes = [
    ['Field', 'Meaning'],
    ['Report group', 'OpenROAD/STA clock group (clk = core, vclk_clk = IO/output virtual clock)'],
    ['Functional path group', 'Data-flow group PG1–PG8 (see flow/analysis/path_groups.md)'],
    ['Launch clk latency', 'Clock network delay from clock source to launch register CLK pin'],
    ['Data-path delay', 'data arrival time − launch clock latency (combinational logic delay)'],
    ['Data-path cells (logic levels)', 'Number of cell instances in the data path (launch Q → endpoint)'],
    ['Capture clk latency', 'Clock network delay from clock source to capture register CLK pin'],
    ['Clock skew', 'capture clk latency − launch clk latency'],
    ['Data required time', 'capture edge + capture latency − setup/output delay (setup)'],
    ['Slack / WNS', 'data required time − data arrival time (setup); arrival − required (hold)'],
    ['Clock period', f'{PERIOD} ns (142.86 MHz); source flow/timing/constraint_7ns.sdc'],
    ['Abbreviations', 'see flow/analysis/glossary.md'],
]
for row in notes:
    ws4.append(row)
ws4['A1'].font = HDR
ws4['A1'].fill = HFILL
ws4['B1'].font = HDR
ws4['B1'].fill = HFILL
ws4.column_dimensions['A'].width = 30
ws4.column_dimensions['B'].width = 80

wb.save(OUT)
print('wrote', OUT, 'with', len(rows), 'paths')
for r in rows:
    print(f"{r['type']:>3} {r['group']:>9} {r['fgroup']:>26} {str(r['start']):>30} -> {str(r['end']):>28} "
          f"launch={r['launch']} data={r['data']} lv={r['levels']} cap={r['capture']} skew={r['skew']} slack={r['slack']}")
