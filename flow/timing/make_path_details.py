#!/usr/bin/env python3
"""Build a baseline workbook with one sheet per worst-5 SETUP path for
PG2 (ALU/write-back, group clk) and PG8 (result-out, group vclk_clk),
broken down into per-cell and per-net (wire) delay."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = '/mnt/d/opencode/details.rpt'
OUT = '/mnt/d/opencode/lunarch/flow/timing/riscv32i_worst10_path_details.xlsx'

text = open(SRC, encoding='utf-8', errors='ignore').read()
blocks = re.split(r'\n(?=Startpoint:)', text)

paths = []
for b in blocks:
    b = b.strip('\n')
    if not b.startswith('Startpoint:'):
        continue
    lines = b.splitlines()
    start = lines[0].split('Startpoint:', 1)[1].strip().split(' (')[0]
    end = re.search(r'Endpoint:\s*(\S+)', b).group(1)
    grp = re.search(r'Path Group:\s*(\S+)', b).group(1)
    arr = float(re.search(r'^\s*(-?[\d.]+)\s+data arrival time', b, re.M).group(1))
    req = float(re.search(r'^\s*(-?[\d.]+)\s+data required time', b, re.M).group(1))
    sl = float(re.search(r'^\s*(-?[\d.]+)\s+slack', b, re.M).group(1))
    start_i = next((i for i, l in enumerate(lines) if '/Q (' in l), None)
    end_i = next((i for i, l in enumerate(lines) if 'data arrival time' in l), len(lines))
    stages = []
    pending = None
    if start_i is not None:
        for l in lines[start_i + 1:end_i]:
            mo = re.match(r'^\s*\d+\s+[\d.]+\s+[\d.]+\s+([\d.]+)\s+([\d.]+)\s+[\^v]\s+(\S+)\s+\((\S+)\)', l)
            if mo:
                cd, t, pin, ct = mo.groups()
                stages.append(dict(pin=pin, ct=ct, cell=float(cd), wire=pending, time=float(t)))
                pending = None
                continue
            mi = re.match(r'^\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+[\^v]\s+(\S+)\s+\((\S+)\)', l)
            if mi:
                pending = float(mi.group(2))
    paths.append(dict(start=start, end=end, group=grp, arr=arr, req=req, slack=sl, stages=stages))

pg2 = [p for p in paths if p['group'] == 'clk']
pg8 = [p for p in paths if p['group'] == 'vclk_clk']

HDR = Font(bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor='3B6FB0')
HF2 = PatternFill('solid', fgColor='7A5CA8')
thin = Side(style='thin', color='BBBBBB')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
TITLE = Font(bold=True, size=13)
SUB = Font(bold=True)
ALT = PatternFill('solid', fgColor='F2F6FC')

wb = Workbook()
idx = wb.active
idx.title = 'Index'


def add_sheet(name, label, p, fill):
    ws = wb.create_sheet(name)
    ws['A1'] = f'{label}:  {p["start"]}  ->  {p["end"]}'
    ws['A1'].font = TITLE
    lv = len(p['stages'])
    ctot = sum(s['cell'] for s in p['stages'])
    wtot = sum(s['wire'] for s in p['stages'] if s['wire'] is not None)
    meta = [
        ('Report (STA) group', p['group']),
        ('Data arrival time (ns)', p['arr']),
        ('Data required time (ns)', p['req']),
        ('Slack / WNS (ns)', p['slack']),
        ('Logic levels (cells)', lv),
        ('Total cell delay (ns)', round(ctot, 4)),
        ('Total wire delay (ns)', round(wtot, 4)),
        ('Cell + wire (ns)', round(ctot + wtot, 4)),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        ws.cell(row=i, column=1, value=k).font = SUB
        ws.cell(row=i, column=2, value=v)
    hdr = ['#', 'Stage (instance)', 'Cell type', 'Cell delay (ns)', 'Wire delay (ns)', 'Time (ns)']
    ws.append([])
    ws.append(hdr)
    hr = ws.max_row
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = HDR
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
    for i, s in enumerate(p['stages'], 1):
        ws.append([i, s['pin'], s['ct'], round(s['cell'], 4),
                   round(s['wire'], 4) if s['wire'] is not None else None, round(s['time'], 4)])
    ws.append([])
    ws.append(['', 'TOTAL', '', round(ctot, 4), round(wtot, 4), ''])
    ws.cell(row=ws.max_row, column=2).font = SUB
    for i, w in enumerate([5, 30, 26, 16, 16, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=hr + 1):
        for c in row:
            c.border = BORD
            if isinstance(c.value, float):
                c.number_format = '0.0000'
    return dict(name=name, label=label, start=p['start'], end=p['end'], group=p['group'],
                arrival=p['arr'], required=p['req'], slack=p['slack'], levels=lv,
                cell=round(ctot, 4), wire=round(wtot, 4))


meta_rows = []
for i, p in enumerate(pg2, 1):
    meta_rows.append(add_sheet(f'PG2-{i}', f'PG2 ALU/write-back path {i}', p, HF))
for i, p in enumerate(pg8, 1):
    meta_rows.append(add_sheet(f'PG8-{i}', f'PG8 result-out path {i}', p, HF2))

# Index sheet
idx['A1'] = 'Baseline reference — worst 5 SETUP paths, PG2 (ALU/write-back) and PG8 (result-out)'
idx['A1'].font = TITLE
idx['A2'] = 'Clock period: 7.0 ns (142.86 MHz). Source: flow/timing/6_finish (ODB) report_checks.'
hdr = ['Sheet', 'Functional path group', 'Report group', 'Start point', 'End point',
       'Logic levels', 'Total cell delay (ns)', 'Total wire delay (ns)',
       'Data arrival (ns)', 'Data required (ns)', 'Slack / WNS (ns)']
idx.append([])
idx.append(hdr)
hr = idx.max_row
for c in range(1, len(hdr) + 1):
    cell = idx.cell(row=hr, column=c)
    cell.font = HDR
    cell.fill = HF
    cell.alignment = Alignment(horizontal='center')
for m in meta_rows:
    idx.append([m['name'], m['label'].split(' path')[0], m['group'], m['start'], m['end'],
                m['levels'], m['cell'], m['wire'], m['arrival'], m['required'], m['slack']])
for i, w in enumerate([8, 22, 12, 32, 22, 12, 18, 18, 14, 15, 14], 1):
    idx.column_dimensions[get_column_letter(i)].width = w
idx.freeze_panes = 'A4'

# Definitions sheet
d = wb.create_sheet('Definitions')
defs = [
    ['Field', 'Meaning'],
    ['Stage (instance)', 'Cell instance name along the path (output pin)'],
    ['Cell delay (ns)', 'Delay through that cell (its timing arc)'],
    ['Wire delay (ns)', 'Interconnect (net RC) delay feeding that cell'],
    ['Time (ns)', 'Cumulative arrival time at that cell output'],
    ['Logic levels', 'Number of cells in the data path'],
    ['PG2 ALU/write-back', 'regfile read -> mux/ALU -> mux -> regfile write (group clk)'],
    ['PG8 result-out', 'regfile read -> ... -> aluout output port (group vclk_clk)'],
    ['Cell delay total', 'Sum of all cell delays'],
    ['Wire delay total', 'Sum of all interconnect delays'],
]
for r in defs:
    d.append(r)
d['A1'].font = HDR
d['A1'].fill = HF
d['B1'].font = HDR
d['B1'].fill = HF
d.column_dimensions['A'].width = 26
d.column_dimensions['B'].width = 72

wb.save(OUT)
print('wrote', OUT)
for m in meta_rows:
    print(f"{m['name']:6} {m['group']:9} lv={m['levels']:2} cell={m['cell']:.4f} wire={m['wire']:.4f} slack={m['slack']}")
